"""Category budgets with alerts, and the accounting export."""
import csv
import io
from datetime import timedelta
from decimal import Decimal

import pytest
from django.core import mail
from django.utils import timezone
from openpyxl import load_workbook

from apps.categories.budgets import budget_status, check_budget_alert
from apps.categories.models import CategoryBudget
from apps.expenses.models import Expense


@pytest.fixture
def food(system_categories):
    return system_categories[0]


@pytest.fixture
def budget(user, food):
    return CategoryBudget.objects.create(user=user, category=food, monthly_amount=Decimal('100000'))


def _spend(user, category, supplier, amount, when=None):
    return Expense.objects.create(
        user=user, category=category, supplier=supplier, description='Mercado',
        amount=Decimal(amount), currency='COP', incurred_at=when or timezone.now(),
    )


class TestBudgetStatus:
    def test_only_counts_current_month(self, user, food, supplier, budget):
        _spend(user, food, supplier, '30000')
        _spend(user, food, supplier, '90000', timezone.now() - timedelta(days=62))
        status = budget_status(budget)
        assert status['spent'] == Decimal('30000')
        assert status['percentage'] == Decimal('30.0')
        assert status['state'] == 'ok'

    def test_states(self, user, food, supplier, budget):
        _spend(user, food, supplier, '85000')
        assert budget_status(budget)['state'] == 'warning'
        _spend(user, food, supplier, '20000')
        assert budget_status(budget)['state'] == 'exceeded'


class TestBudgetAlerts:
    def test_warning_then_exceeded_each_sent_once(self, user, food, supplier, budget):
        _spend(user, food, supplier, '85000')
        assert check_budget_alert(user, food) == 'warning'
        assert check_budget_alert(user, food) is None
        _spend(user, food, supplier, '20000')
        assert check_budget_alert(user, food) == 'exceeded'
        assert check_budget_alert(user, food) is None
        assert len(mail.outbox) == 2
        assert 'Superaste' in mail.outbox[1].subject

    def test_api_expense_triggers_alert(self, auth_client, food, supplier, budget):
        response = auth_client.post('/api/v1/expenses/', {
            'category': food.pk, 'supplier': supplier.pk, 'description': 'Mercado grande',
            'amount': '120000', 'currency': 'COP', 'payment_method': 'card',
            'incurred_at': timezone.now().isoformat(),
        }, format='json')
        assert response.status_code == 201
        assert len(mail.outbox) == 1

    def test_old_expense_does_not_alert(self, auth_client, food, supplier, budget):
        auth_client.post('/api/v1/expenses/', {
            'category': food.pk, 'supplier': supplier.pk, 'description': 'Factura vieja',
            'amount': '500000', 'currency': 'COP', 'payment_method': 'card',
            'incurred_at': (timezone.now() - timedelta(days=400)).isoformat(),
        }, format='json')
        assert len(mail.outbox) == 0


class TestBudgetApi:
    def test_create_is_upsert_and_lists_status(self, auth_client, food):
        first = auth_client.post('/api/v1/budgets/', {'category': food.pk, 'monthly_amount': '50000'}, format='json')
        assert first.status_code == 201
        second = auth_client.post('/api/v1/budgets/', {'category': food.pk, 'monthly_amount': '70000'}, format='json')
        assert second.status_code == 200
        listing = auth_client.get('/api/v1/budgets/')
        assert len(listing.data['data']) == 1
        assert listing.data['data'][0]['monthly_amount'] == '70000.00'
        assert listing.data['data'][0]['status']['state'] == 'ok'

    def test_rejects_other_users_category(self, auth_client, user2):
        from apps.categories.models import Category
        foreign = Category.objects.create(user=user2, name='Ajena', scope=Category.SCOPE_USER)
        response = auth_client.post('/api/v1/budgets/', {'category': foreign.pk, 'monthly_amount': '1000'}, format='json')
        assert response.status_code == 400

    def test_rejects_zero(self, auth_client, food):
        response = auth_client.post('/api/v1/budgets/', {'category': food.pk, 'monthly_amount': '0'}, format='json')
        assert response.status_code == 400


class TestExport:
    def test_csv(self, auth_client, user, food, supplier):
        _spend(user, food, supplier, '10000')
        response = auth_client.get('/api/v1/expenses/export/?file_format=csv')
        assert response.status_code == 200
        assert 'attachment' in response['Content-Disposition']
        text = response.content.decode('utf-8-sig')
        rows = list(csv.reader(io.StringIO(text)))
        assert rows[0][0] == 'Fecha'
        assert rows[1][1] == supplier.name
        assert Decimal(rows[1][11]) == Decimal('10000')

    def test_xlsx_with_summaries_and_date_filter(self, auth_client, user, food, supplier):
        _spend(user, food, supplier, '10000')
        _spend(user, food, supplier, '5000', timezone.now() - timedelta(days=90))
        start = (timezone.localdate() - timedelta(days=10)).isoformat()
        response = auth_client.get(f'/api/v1/expenses/export/?file_format=xlsx&start_date={start}')
        assert response.status_code == 200
        workbook = load_workbook(io.BytesIO(response.content))
        assert workbook.sheetnames == ['Gastos', 'Por categoría', 'Por proveedor']
        sheet = workbook['Gastos']
        assert sheet.max_row == 3  # header, one expense, total row
        assert workbook['Por categoría'].cell(row=2, column=1).value == food.name

    def test_bad_format(self, auth_client):
        assert auth_client.get('/api/v1/expenses/export/?file_format=pdf').status_code == 400

    def test_only_own_expenses(self, auth_client, user2, food):
        Expense.objects.create(
            user=user2, category=food, description='Ajeno', amount=Decimal('1'),
            incurred_at=timezone.now(),
        )
        text = auth_client.get('/api/v1/expenses/export/?file_format=csv').content.decode('utf-8-sig')
        assert 'Ajeno' not in text
