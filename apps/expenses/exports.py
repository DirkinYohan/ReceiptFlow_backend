"""Accounting export of expenses: CSV for other software, Excel for people."""
from __future__ import annotations

import csv
import io
from collections import defaultdict
from decimal import Decimal

from django.utils import timezone

from .models import Expense

COLUMNS = [
    ('Fecha', 12),
    ('Proveedor', 32),
    ('NIT', 16),
    ('N° factura', 16),
    ('CUFE', 30),
    ('Categoría', 16),
    ('Descripción', 36),
    ('Base', 14),
    ('Impuestos', 14),
    ('Descuento', 12),
    ('Propina', 12),
    ('Total', 14),
    ('Moneda', 8),
    ('Método de pago', 14),
    ('Pagado', 8),
    ('Recurrente', 12),
    ('ID gasto', 9),
    ('ID documento', 12),
]
MONEY_COLUMNS = {'Base', 'Impuestos', 'Descuento', 'Propina', 'Total'}


def export_queryset(user, start=None, end=None, category_id=None):
    qs = Expense.objects.filter(user=user).select_related('supplier', 'category', 'document')
    if start:
        qs = qs.filter(incurred_at__date__gte=start)
    if end:
        qs = qs.filter(incurred_at__date__lte=end)
    if category_id:
        qs = qs.filter(category_id=category_id)
    return qs.order_by('incurred_at', 'id')


def _row(expense: Expense) -> list:
    document = expense.document
    structured = (document.structured_data or {}) if document else {}
    recurrence = dict(Expense.RECURRENCE_CHOICES).get(expense.recurrence_frequency, '')
    return [
        timezone.localdate(expense.incurred_at),
        expense.supplier.name if expense.supplier else '',
        expense.supplier.tax_id if expense.supplier else '',
        document.document_number if document else '',
        structured.get('cufe') or '',
        expense.category.name,
        expense.description,
        expense.amount,
        expense.tax_amount,
        expense.discount_amount,
        expense.tip_amount,
        expense.total_amount,
        expense.currency,
        expense.get_payment_method_display(),
        ('Sí' if document.is_paid else 'No') if document else '',
        recurrence if expense.is_recurring else '',
        expense.pk,
        document.pk if document else '',
    ]


def build_csv(expenses) -> bytes:
    buffer = io.StringIO()
    writer = csv.writer(buffer)
    writer.writerow([name for name, _ in COLUMNS])
    for expense in expenses:
        writer.writerow([
            value.isoformat() if hasattr(value, 'isoformat') else value
            for value in _row(expense)
        ])
    # BOM so Excel opens the accents correctly.
    return ('﻿' + buffer.getvalue()).encode('utf-8')


def _summary(expenses, key) -> list[tuple[str, int, Decimal]]:
    grouped: dict = defaultdict(lambda: [0, Decimal('0')])
    for expense in expenses:
        name = key(expense)
        grouped[name][0] += 1
        grouped[name][1] += expense.total_amount
    return sorted(((name, count, total) for name, (count, total) in grouped.items()), key=lambda r: -r[2])


def build_xlsx(expenses, title: str) -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    expenses = list(expenses)
    header_font = Font(bold=True, color='FFFFFF')
    header_fill = PatternFill('solid', fgColor='047857')
    money_format = '#,##0.00'

    def write_header(sheet, names_widths):
        sheet.append([name for name, _ in names_widths])
        for idx, (_, width) in enumerate(names_widths, start=1):
            cell = sheet.cell(row=1, column=idx)
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = Alignment(vertical='center')
            sheet.column_dimensions[get_column_letter(idx)].width = width
        sheet.freeze_panes = 'A2'

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = 'Gastos'
    write_header(sheet, COLUMNS)
    money_idx = [i for i, (name, _) in enumerate(COLUMNS, start=1) if name in MONEY_COLUMNS]
    for expense in expenses:
        sheet.append(_row(expense))
        row = sheet.max_row
        sheet.cell(row=row, column=1).number_format = 'DD/MM/YYYY'
        for idx in money_idx:
            sheet.cell(row=row, column=idx).number_format = money_format
    if expenses:
        total_row = sheet.max_row + 1
        sheet.cell(row=total_row, column=1, value='TOTAL').font = Font(bold=True)
        for idx in money_idx:
            letter = get_column_letter(idx)
            cell = sheet.cell(row=total_row, column=idx, value=f'=SUM({letter}2:{letter}{total_row - 1})')
            cell.font = Font(bold=True)
            cell.number_format = money_format
        sheet.auto_filter.ref = f'A1:{get_column_letter(len(COLUMNS))}{total_row - 1}'

    for sheet_title, key in (
        ('Por categoría', lambda e: e.category.name),
        ('Por proveedor', lambda e: e.supplier.name if e.supplier else 'Sin proveedor'),
    ):
        summary = workbook.create_sheet(sheet_title)
        write_header(summary, [(sheet_title.replace('Por ', '').capitalize(), 32), ('Gastos', 10), ('Total', 16)])
        for name, count, total in _summary(expenses, key):
            summary.append([name, count, total])
            summary.cell(row=summary.max_row, column=3).number_format = money_format

    workbook.properties.title = title
    output = io.BytesIO()
    workbook.save(output)
    return output.getvalue()
